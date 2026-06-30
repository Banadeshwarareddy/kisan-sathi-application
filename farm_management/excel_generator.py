from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from datetime import datetime

# Color constants
GREEN_FILL = PatternFill('solid', fgColor='22c55e')
BLUE_FILL = PatternFill('solid', fgColor='3b82f6')
ORANGE_FILL = PatternFill('solid', fgColor='f97316')
DARK_FILL = PatternFill('solid', fgColor='1b4332')
GRAY_FILL = PatternFill('solid', fgColor='f3f4f6')
WHITE_FILL = PatternFill('solid', fgColor='ffffff')
LIGHT_GREEN = PatternFill('solid', fgColor='dcfce7')
LIGHT_BLUE = PatternFill('solid', fgColor='dbeafe')

WHITE_FONT = Font(name='Calibri', color='FFFFFF', bold=True)
DARK_FONT = Font(name='Calibri', color='1a1a1a')
BOLD_FONT = Font(name='Calibri', bold=True)
GREEN_FONT = Font(name='Calibri', color='16a34a', bold=True)
RED_FONT = Font(name='Calibri', color='dc2626', bold=True)

THIN_BORDER = Border(
    left=Side(style='thin', color='e5e7eb'),
    right=Side(style='thin', color='e5e7eb'),
    top=Side(style='thin', color='e5e7eb'),
    bottom=Side(style='thin', color='e5e7eb'),
)

CENTER = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center')
RIGHT = Alignment(horizontal='right', vertical='center')


def auto_fit_columns(ws, min_width=10, max_width=40):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, min_width), max_width)


def generate_expense_excel(expenses, farmer) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Expenses'

    farmer_name = farmer.get_full_name() or farmer.username
    total = sum(e.amount for e in expenses)

    # Title
    ws.merge_cells('A1:E1')
    ws['A1'] = '🌾 KISAN SATHI - EXPENSE REPORT'
    ws['A1'].fill = DARK_FILL
    ws['A1'].font = Font(name='Calibri', color='FFFFFF', bold=True, size=14)
    ws['A1'].alignment = CENTER
    ws.row_dimensions[1].height = 35

    ws.merge_cells('A2:E2')
    ws['A2'] = f'Farmer: {farmer_name} | Generated: {datetime.now().strftime("%d %b %Y %I:%M %p")}'
    ws['A2'].fill = PatternFill('solid', fgColor='2d6a4f')
    ws['A2'].font = Font(name='Calibri', color='FFFFFF', size=9)
    ws['A2'].alignment = CENTER
    ws.row_dimensions[2].height = 18

    # Summary cards row
    ws['A4'] = 'Total Expenses'
    ws['B4'] = f'₹{total:,.2f}'
    ws['C4'] = 'Total Records'
    ws['D4'] = expenses.count()
    ws['E4'] = f'Report: {datetime.now().strftime("%d %b %Y")}'

    for col in range(1, 6):
        cell = ws.cell(row=4, column=col)
        cell.fill = LIGHT_BLUE if col in [1,2] else LIGHT_GREEN
        cell.font = BOLD_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER
    ws.row_dimensions[4].height = 22

    # Column headers
    headers = ['#', 'Category', 'Amount (₹)', 'Date', 'Notes']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col, value=header)
        cell.fill = BLUE_FILL
        cell.font = WHITE_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER
    ws.row_dimensions[6].height = 20

    # Data rows
    for row_num, expense in enumerate(expenses, 7):
        row_data = [
            row_num - 6,
            expense.get_category_display(),
            float(expense.amount),
            expense.date.strftime('%d %b %Y'),
            expense.notes,
        ]
        fill = WHITE_FILL if row_num % 2 == 0 else GRAY_FILL
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.fill = fill
            cell.font = DARK_FONT
            cell.border = THIN_BORDER
            cell.alignment = RIGHT if col == 3 else (CENTER if col in [1, 4] else LEFT)
            if col == 3:
                cell.number_format = '₹#,##0.00'

    # Total row
    total_row = expenses.count() + 7
    total_data = ['', 'TOTAL', float(total), '', '']
    for col, value in enumerate(total_data, 1):
        cell = ws.cell(row=total_row, column=col, value=value)
        cell.fill = LIGHT_BLUE
        cell.font = Font(name='Calibri', bold=True, color='1d4ed8', size=11)
        cell.border = THIN_BORDER
        cell.alignment = RIGHT if col == 3 else CENTER
        if col == 3:
            cell.number_format = '₹#,##0.00'

    auto_fit_columns(ws)
    ws.column_dimensions['E'].width = 40
    ws.freeze_panes = 'A7'

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_income_excel(incomes, farmer) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Income'

    farmer_name = farmer.get_full_name() or farmer.username
    total = sum(i.total_amount or 0 for i in incomes)

    # Title
    ws.merge_cells('A1:H1')
    ws['A1'] = '🌾 KISAN SATHI - INCOME REPORT'
    ws['A1'].fill = DARK_FILL
    ws['A1'].font = Font(name='Calibri', color='FFFFFF', bold=True, size=14)
    ws['A1'].alignment = CENTER
    ws.row_dimensions[1].height = 35

    ws.merge_cells('A2:H2')
    ws['A2'] = f'Farmer: {farmer_name} | Generated: {datetime.now().strftime("%d %b %Y %I:%M %p")}'
    ws['A2'].fill = PatternFill('solid', fgColor='2d6a4f')
    ws['A2'].font = Font(name='Calibri', color='FFFFFF', size=9)
    ws['A2'].alignment = CENTER

    # Summary
    ws['A4'] = 'Total Income'
    ws['B4'] = f'₹{total:,.2f}'
    ws['C4'] = 'Records'
    ws['D4'] = incomes.count()
    ws['E4'] = 'Received'
    ws['F4'] = sum(1 for i in incomes if i.payment_status=='received')
    ws['G4'] = 'Pending'
    ws['H4'] = sum(1 for i in incomes if i.payment_status=='pending')

    for col in range(1, 9):
        cell = ws.cell(row=4, column=col)
        cell.fill = LIGHT_GREEN if col in [1,2] else LIGHT_BLUE
        cell.font = BOLD_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    # Headers
    headers = ['#', 'Crop', 'Quantity', 'Unit', 'Rate/Unit (₹)', 'Total (₹)', 'Buyer', 'Date', 'Status']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col, value=header)
        cell.fill = GREEN_FILL
        cell.font = WHITE_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    # Data
    for row_num, inc in enumerate(incomes, 7):
        row_data = [
            row_num - 6,
            inc.crop,
            float(inc.quantity) if inc.quantity else '-',
            inc.get_unit_display(),
            float(inc.rate_per_unit) if inc.rate_per_unit else '-',
            float(inc.total_amount) if inc.total_amount else 0,
            inc.buyer_name or '-',
            inc.sale_date.strftime('%d %b %Y'),
            inc.payment_status.title(),
        ]
        fill = WHITE_FILL if row_num % 2 == 0 else GRAY_FILL
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.fill = fill
            cell.font = DARK_FONT
            cell.border = THIN_BORDER
            cell.alignment = RIGHT if col in [5,6] else CENTER
            if col in [5,6] and isinstance(value, float):
                cell.number_format = '₹#,##0.00'

        # Color payment status
        status_cell = ws.cell(row=row_num, column=9)
        if inc.payment_status == 'received':
            status_cell.font = GREEN_FONT
        elif inc.payment_status == 'pending':
            status_cell.font = RED_FONT

    # Total row
    total_row = incomes.count() + 7
    for col in range(1, 10):
        cell = ws.cell(row=total_row, column=col)
        cell.fill = LIGHT_GREEN
        cell.font = Font(name='Calibri', bold=True, color='16a34a', size=11)
        cell.border = THIN_BORDER
        cell.alignment = CENTER

    ws.cell(row=total_row, column=2, value='TOTAL')
    ws.cell(row=total_row, column=6, value=float(total)).number_format = '₹#,##0.00'

    auto_fit_columns(ws)
    ws.freeze_panes = 'A7'

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_full_report_excel(farmer) -> BytesIO:
    """Single-sheet complete farm report Excel matching the PDF layout"""
    from .models import Expense, Income

    wb = Workbook()
    ws = wb.active
    ws.title = 'Farm Report'
    ws.views.sheetView[0].showGridLines = True

    expenses = Expense.objects.filter(farmer=farmer, is_deleted=False)
    incomes = Income.objects.filter(farmer=farmer, is_deleted=False)

    total_exp = sum(e.amount for e in expenses)
    total_inc = sum(i.total_amount or 0 for i in incomes)
    net = total_inc - total_exp

    farmer_name = farmer.get_full_name() or farmer.username

    # Colors matching the PDF styling
    green_header_fill = PatternFill('solid', fgColor='16a34a')
    blue_header_fill = PatternFill('solid', fgColor='3b82f6')
    dark_fill = PatternFill('solid', fgColor='1b4332')
    sub_fill = PatternFill('solid', fgColor='2d6a4f')
    
    light_green = PatternFill('solid', fgColor='dcfce7')
    light_orange = PatternFill('solid', fgColor='fff7ed')
    light_blue = PatternFill('solid', fgColor='dbeafe')
    
    white_fill = PatternFill('solid', fgColor='ffffff')
    gray_fill = PatternFill('solid', fgColor='f9fafb')

    # Font styles
    white_font_bold = Font(name='Calibri', color='FFFFFF', bold=True, size=11)
    dark_font = Font(name='Calibri', color='1a1a1a', size=10)
    bold_font = Font(name='Calibri', bold=True, size=10)

    def style_range(cell_range, border=None, fill=None, font=None, alignment=None):
        for row in ws[cell_range]:
            for cell in row:
                if border:
                    cell.border = border
                if fill:
                    cell.fill = fill
                if font:
                    cell.font = font
                if alignment:
                    cell.alignment = alignment

    # 1. Header Banner
    ws.merge_cells('A1:E1')
    ws['A1'] = 'KISAN SATHI'
    style_range('A1:E1', fill=dark_fill, font=Font(name='Calibri', color='FFFFFF', bold=True, size=16), alignment=LEFT)
    ws.row_dimensions[1].height = 30

    ws.merge_cells('A2:E2')
    ws['A2'] = f'Complete Farm Report | Net Profit: ₹{net:,.2f}'
    style_range('A2:E2', fill=sub_fill, font=Font(name='Calibri', color='FFFFFF', bold=True, size=11), alignment=LEFT)
    ws.row_dimensions[2].height = 20

    ws.merge_cells('A3:E3')
    ws['A3'] = f'Farmer: {farmer_name} | Generated: {datetime.now().strftime("%d %b %Y %I:%M %p")}'
    style_range('A3:E3', fill=sub_fill, font=Font(name='Calibri', color='FFFFFF', size=9), alignment=LEFT)
    ws.row_dimensions[3].height = 18

    # Spacer row
    ws.row_dimensions[4].height = 12

    # 2. Financial Summary Table
    ws.merge_cells('A5:E5')
    ws['A5'] = 'Financial Summary'
    style_range('A5:E5', fill=green_header_fill, font=white_font_bold, alignment=LEFT)
    ws.row_dimensions[5].height = 22

    ws.cell(row=6, column=1, value='Total Income').font = bold_font
    ws.cell(row=6, column=1).border = THIN_BORDER
    
    inc_val = ws.cell(row=6, column=2, value=float(total_inc))
    inc_val.font = bold_font
    inc_val.fill = light_green
    inc_val.number_format = '₹#,##0.00'
    inc_val.border = THIN_BORDER
    
    ws.cell(row=6, column=3, value='Total Expenses').font = bold_font
    ws.cell(row=6, column=3).border = THIN_BORDER
    
    exp_val = ws.cell(row=6, column=4, value=float(total_exp))
    exp_val.font = bold_font
    exp_val.fill = light_orange
    exp_val.number_format = '₹#,##0.00'
    exp_val.border = THIN_BORDER

    ws.cell(row=6, column=5).border = THIN_BORDER

    ws.cell(row=7, column=1, value='Net Profit/Loss').font = bold_font
    ws.cell(row=7, column=1).border = THIN_BORDER
    
    net_val = ws.cell(row=7, column=2, value=float(net))
    net_val.font = bold_font
    net_val.fill = light_green if net >= 0 else PatternFill('solid', fgColor='fee2e2')
    net_val.number_format = '₹#,##0.00'
    net_val.border = THIN_BORDER
    
    ws.cell(row=7, column=3).border = THIN_BORDER
    ws.cell(row=7, column=4).border = THIN_BORDER
    ws.cell(row=7, column=5).border = THIN_BORDER

    ws.row_dimensions[6].height = 18
    ws.row_dimensions[7].height = 18

    # Spacer row
    ws.row_dimensions[8].height = 12

    # 3. Expenses Section
    current_row = 9
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    ws.cell(row=current_row, column=1, value='Expenses')
    style_range(f'A{current_row}:E{current_row}', fill=blue_header_fill, font=white_font_bold, alignment=LEFT)
    ws.row_dimensions[current_row].height = 22

    # Headers
    current_row += 1
    ws.merge_cells(start_row=current_row, start_column=4, end_row=current_row, end_column=5)
    ws.cell(row=current_row, column=1, value='Category')
    ws.cell(row=current_row, column=2, value='Amount')
    ws.cell(row=current_row, column=3, value='Date')
    ws.cell(row=current_row, column=4, value='Notes')
    style_range(f'A{current_row}:E{current_row}', border=THIN_BORDER, fill=blue_header_fill, font=white_font_bold, alignment=CENTER)
    ws.row_dimensions[current_row].height = 20

    # Data rows
    for idx, e in enumerate(expenses, 1):
        current_row += 1
        fill = white_fill if idx % 2 == 0 else gray_fill
        ws.merge_cells(start_row=current_row, start_column=4, end_row=current_row, end_column=5)
        
        ws.cell(row=current_row, column=1, value=e.get_category_display())
        ws.cell(row=current_row, column=2, value=float(e.amount)).number_format = '₹#,##0.00'
        ws.cell(row=current_row, column=3, value=e.date.strftime('%d %b %Y'))
        ws.cell(row=current_row, column=4, value=e.notes)
        
        style_range(f'A{current_row}:E{current_row}', border=THIN_BORDER, fill=fill, font=dark_font)
        ws.cell(row=current_row, column=1).alignment = LEFT
        ws.cell(row=current_row, column=2).alignment = RIGHT
        ws.cell(row=current_row, column=3).alignment = CENTER
        ws.cell(row=current_row, column=4).alignment = LEFT

    # Total Row
    current_row += 1
    ws.merge_cells(start_row=current_row, start_column=4, end_row=current_row, end_column=5)
    ws.cell(row=current_row, column=1, value='TOTAL')
    ws.cell(row=current_row, column=2, value=float(total_exp)).number_format = '₹#,##0.00'
    style_range(f'A{current_row}:E{current_row}', border=THIN_BORDER, fill=light_blue, font=bold_font)
    ws.cell(row=current_row, column=1).alignment = LEFT
    ws.cell(row=current_row, column=2).alignment = RIGHT

    # Spacer row
    current_row += 1
    ws.row_dimensions[current_row].height = 12

    # 4. Income Section
    current_row += 1
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    ws.cell(row=current_row, column=1, value='Income')
    style_range(f'A{current_row}:E{current_row}', fill=green_header_fill, font=white_font_bold, alignment=LEFT)
    ws.row_dimensions[current_row].height = 22

    # Headers
    current_row += 1
    ws.cell(row=current_row, column=1, value='Crop')
    ws.cell(row=current_row, column=2, value='Quantity')
    ws.cell(row=current_row, column=3, value='Total')
    ws.cell(row=current_row, column=4, value='Buyer')
    ws.cell(row=current_row, column=5, value='Date')
    style_range(f'A{current_row}:E{current_row}', border=THIN_BORDER, fill=green_header_fill, font=white_font_bold, alignment=CENTER)
    ws.row_dimensions[current_row].height = 20

    # Data rows
    for idx, inc in enumerate(incomes, 1):
        current_row += 1
        fill = white_fill if idx % 2 == 0 else gray_fill
        
        ws.cell(row=current_row, column=1, value=inc.crop)
        qty_str = f"{inc.quantity} {inc.get_unit_display()}" if inc.quantity else "-"
        ws.cell(row=current_row, column=2, value=qty_str)
        ws.cell(row=current_row, column=3, value=float(inc.total_amount) if inc.total_amount else 0.0).number_format = '₹#,##0.00'
        ws.cell(row=current_row, column=4, value=inc.buyer_name or '-')
        ws.cell(row=current_row, column=5, value=inc.sale_date.strftime('%d %b %Y'))
        
        style_range(f'A{current_row}:E{current_row}', border=THIN_BORDER, fill=fill, font=dark_font)
        ws.cell(row=current_row, column=1).alignment = LEFT
        ws.cell(row=current_row, column=2).alignment = CENTER
        ws.cell(row=current_row, column=3).alignment = RIGHT
        ws.cell(row=current_row, column=4).alignment = LEFT
        ws.cell(row=current_row, column=5).alignment = CENTER

    # Total Row
    current_row += 1
    ws.cell(row=current_row, column=1, value='TOTAL')
    ws.cell(row=current_row, column=3, value=float(total_inc)).number_format = '₹#,##0.00'
    style_range(f'A{current_row}:E{current_row}', border=THIN_BORDER, fill=light_green, font=bold_font)
    ws.cell(row=current_row, column=1).alignment = LEFT
    ws.cell(row=current_row, column=3).alignment = RIGHT

    # Formatting adjustments
    auto_fit_columns(ws)
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 20

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
